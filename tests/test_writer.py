from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from src.common.config import Rules
from src.common.country import CountryResolver
from src.common.models import OrderRow
from src.reader.orders import group_by_order
from src.transformer.headers import load_template_headers
from src.transformer.transform import transform_groups
from src.writer.output import (
    extract_order_no,
    filter_new_rows,
    write_output,
)


def test_extract_order_no():
    assert extract_order_no("E-0606-4069437284") == "4069437284"
    assert extract_order_no("C-1231-999") == "999"


def test_append_dedup_by_order_no(tmp_path):
    rules = Rules.load()
    resolver = CountryResolver(rules)
    template_headers = load_template_headers(rules.template_path)

    def make_row(order_no: str) -> OrderRow:
        return OrderRow(
            order_no=order_no,
            style="T恤",
            quantity=1,
            contact="123",
            recipient="User",
            country="Canada",
            city="Toronto",
            state="ON",
            zip_code="M5V",
            address="addr",
            source_row=1,
        )

    out_path = tmp_path / "out.xlsx"
    batch1 = transform_groups([[make_row("1001")]], rules, resolver, date(2026, 6, 6))
    written1, skipped1 = write_output(
        out_path, batch1, rules, template_headers, append=False
    )
    assert written1 == 1
    assert skipped1 == []

    batch2 = transform_groups(
        [[make_row("1001")], [make_row("1002")]], rules, resolver, date(2026, 6, 6)
    )
    written2, skipped2 = write_output(
        out_path, batch2, rules, template_headers, append=True
    )
    assert written2 == 1
    assert skipped2 == ["1001"]

    wb = load_workbook(out_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    order_col = next(i for i, h in enumerate(headers) if str(h).startswith("用户订单号"))
    order_ids = []
    for r in range(2, ws.max_row + 1):
        val = ws.cell(r, order_col + 1).value
        if val:
            order_ids.append(extract_order_no(str(val)))
    wb.close()
    assert order_ids == ["1001", "1002"]


def test_filter_new_rows_within_batch():
    rules = Rules.load()
    resolver = CountryResolver(rules)
    rows = transform_groups(
        [
            [
                OrderRow(
                    order_no="dup",
                    style="T恤",
                    quantity=1,
                    contact=None,
                    recipient="A",
                    country="Canada",
                    city="c",
                    state="s",
                    zip_code="z",
                    address="a",
                    source_row=1,
                )
            ],
            [
                OrderRow(
                    order_no="dup",
                    style="棒球帽",
                    quantity=1,
                    contact=None,
                    recipient="B",
                    country="Canada",
                    city="c",
                    state="s",
                    zip_code="z",
                    address="a",
                    source_row=2,
                )
            ],
        ],
        rules,
        resolver,
        date(2026, 6, 6),
    )
    kept, skipped = filter_new_rows(rows, set())
    assert len(kept) == 1
    assert skipped == ["dup"]


def test_autofit_sets_column_widths(tmp_path):
    rules = Rules.load()
    resolver = CountryResolver(rules)
    groups = [
        [
            OrderRow(
                order_no="1",
                style="T恤",
                quantity=1,
                contact="123",
                recipient="张三 Test User",
                country="Canada",
                city="Toronto",
                state="ON",
                zip_code="M5V",
                address="123 Very Long Street Name For Width",
                source_row=1,
            )
        ]
    ]
    output_rows = transform_groups(groups, rules, resolver, date(2026, 6, 6))
    out_path = tmp_path / "out.xlsx"
    template_headers = load_template_headers(rules.template_path)
    write_output(out_path, output_rows, rules, template_headers, append=False)

    wb = load_workbook(out_path)
    ws = wb.active
    widths = [ws.column_dimensions[chr].width for chr in "ABCDEFGHIJK"]
    wb.close()
    assert any(w > 8.5 for w in widths)
