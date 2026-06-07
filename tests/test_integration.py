from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from src.common.config import Rules
from src.common.country import CountryResolver
from src.common.paths import PROJECT_ROOT
from src.reader.orders import group_by_order, read_orders
from src.transformer.headers import load_template_headers
from src.transformer.transform import transform_groups
from src.writer.output import write_output


def test_process_sample_input(tmp_path):
    rules = Rules.load()
    input_path = PROJECT_ROOT / "file" / "20260606.xlsx"
    rows = read_orders(input_path)
    groups = group_by_order(rows)
    assert len(rows) == 20
    assert len(groups) == 17

    resolver = CountryResolver(rules)
    output_rows = transform_groups(groups, rules, resolver, date(2026, 6, 6))

    out_path = tmp_path / "20260606.xlsx"
    template_headers = load_template_headers(rules.template_path)
    write_output(out_path, output_rows, rules, template_headers, append=False)

    wb = load_workbook(out_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    data_rows = 0
    us_clothing = False
    four_product_cols = False
    for r in range(2, ws.max_row + 1):
        order_id = None
        for c, h in enumerate(headers, 1):
            if h and str(h).startswith("用户订单号"):
                order_id = ws.cell(r, c).value
                break
        if not order_id:
            continue
        data_rows += 1
        if str(order_id).endswith("4069983558"):
            four_product_cols = any(h == "英文申报品名4" for h in headers)
            assert len(headers) == 97
        for c, h in enumerate(headers, 1):
            if h and str(h).startswith("业务类型"):
                bt = ws.cell(r, c).value
                for c2, h2 in enumerate(headers, 1):
                    if h2 and str(h2).startswith("目的地国家代码"):
                        cc = ws.cell(r, c2).value
                        if cc == "US" and bt == "服装专线":
                            us_clothing = True
    wb.close()

    assert data_rows == 17
    assert us_clothing
    assert four_product_cols
