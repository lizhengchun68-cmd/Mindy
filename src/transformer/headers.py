from __future__ import annotations

from openpyxl import load_workbook

from src.common.config import Rules
from src.common.models import OutputRow

BUILTIN_PRODUCT_COUNT = 68  # template cols 1–68 (A–BP)


def load_template_headers(template_path) -> list[str]:
    wb = load_workbook(template_path, read_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    wb.close()
    return headers


def build_headers(max_products: int, template_headers: list[str], rules: Rules) -> list[str]:
    if max_products <= rules.builtin_slots:
        return list(template_headers)

    base = list(template_headers[:BUILTIN_PRODUCT_COUNT])
    for n in range(rules.builtin_slots + 1, max_products + 1):
        for tmpl in rules.extra_header_template:
            base.append(tmpl.replace("{N}", str(n)))
    base.extend(rules.tail_columns)
    return base


def count_product_slots_in_headers(headers: list[str]) -> int:
    return sum(1 for h in headers if h and str(h).startswith("英文申报品名"))


def _find_header(headers: list[str], starts_with: str, *, exclude_contains: str | None = None) -> str | None:
    for h in headers:
        if not h:
            continue
        hs = str(h)
        if exclude_contains and exclude_contains in hs:
            continue
        if hs.startswith(starts_with):
            return hs
    return None


def row_to_values(row: OutputRow, headers: list[str], rules: Rules) -> list:
    fixed = rules.raw["fixed_values"]
    values: dict[str, object] = {}

    field_map: list[tuple[str, object, str | None]] = [
        ("业务类型", row.business_type, None),
        ("用户订单号", row.order_id, None),
        ("收件公司", "", None),
        ("收件联系人", row.contact_name, None),
        ("目的地国家代码", row.country_code, None),
        ("收方州/省", row.state or "", None),
        ("收方城市", row.city or "", "编码"),
        ("收方县/区", "", None),
        ("收方城市编码", "", None),
        ("收方详细地址", row.address or "", None),
        ("收件人门牌号", "", None),
        ("收件邮编", row.zip_code or "", None),
        ("收件人电话", row.phone, None),
        ("收件人手机", row.phone, None),
        ("是否带电", fixed["是否带电"], None),
        ("申报价值币种", fixed["申报价值币种"], None),
    ]

    for prefix, val, exclude in field_map:
        h = _find_header(headers, prefix, exclude_contains=exclude)
        if h:
            values[h] = val

    for idx, product in enumerate(row.products, start=1):
        slot_fields = [
            (f"英文申报品名{idx}", product.en),
            (f"中文申报品名{idx}", product.cn),
            (f"申报价值（单价）{idx}", product.unit_price),
            (f"成本单价{idx}", ""),
            (f"商品网址链接{idx}", ""),
            (f"品名{idx}单位重量", product.unit_weight),
            (f"申报品数量{idx}", product.quantity),
            (f"海关货物编号{idx}", product.hs_code),
            (f"法定数量{idx}", ""),
            (f"第二数量{idx}", ""),
            (f"法定计量单位{idx}", ""),
            (f"第二计量单位{idx}", ""),
        ]
        for starts, val in slot_fields:
            h = _find_header(headers, starts)
            if h:
                values[h] = val

    return [values.get(h, "") for h in headers]
