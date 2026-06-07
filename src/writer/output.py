from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from src.common.config import Rules
from src.common.models import OutputRow
from src.transformer.headers import (
    build_headers,
    count_product_slots_in_headers,
    row_to_values,
)


def get_output_path(output_dir: Path, run_date) -> Path:
    return output_dir / f"{run_date.strftime('%Y%m%d')}.xlsx"


def extract_order_no(order_id: str) -> str:
    """从 B 列用户订单号（如 E-0606-4069437284）提取原始订单编号。"""
    return str(order_id).rsplit("-", 1)[-1].strip()


def _find_order_id_col(headers: list) -> int | None:
    for i, h in enumerate(headers):
        if h and str(h).startswith("用户订单号"):
            return i
    return None


def load_existing_order_nos(path: Path) -> set[str]:
    if not path.exists() or not _has_data_rows(path):
        return set()

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col_idx = _find_order_id_col(headers)
        if col_idx is None:
            return set()

        order_nos: set[str] = set()
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, col_idx + 1).value
            if val is not None and str(val).strip():
                order_nos.add(extract_order_no(str(val)))
        return order_nos
    finally:
        wb.close()


def filter_new_rows(
    new_rows: list[OutputRow], existing_order_nos: set[str]
) -> tuple[list[OutputRow], list[str]]:
    """按原始订单编号去重，跳过已存在或本批重复的订单。"""
    seen = set(existing_order_nos)
    kept: list[OutputRow] = []
    skipped: list[str] = []
    for row in new_rows:
        order_no = extract_order_no(row.order_id)
        if order_no in seen:
            skipped.append(order_no)
            continue
        seen.add(order_no)
        kept.append(row)
    return kept, skipped


def find_last_data_row(ws, marker_prefix: str = "用户订单号") -> int:
    last = 1
    marker_col = None
    for c in range(1, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h and str(h).startswith(marker_prefix):
            marker_col = c
            break
    if not marker_col:
        return 1

    for r in range(2, ws.max_row + 1):
        val = ws.cell(r, marker_col).value
        if val is not None and str(val).strip():
            last = r
    return last


def _display_width(text: str) -> float:
    """Estimate Excel column width; CJK chars count wider than ASCII."""
    width = 0.0
    for ch in text:
        width += 2.0 if ord(ch) > 127 else 1.0
    return width


def autofit_column_widths(ws, *, min_width: float = 8, max_width: float = 48, padding: float = 2) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        longest = 0.0
        for row_idx in range(1, ws.max_row + 1):
            val = ws.cell(row_idx, col_idx).value
            if val is None:
                continue
            longest = max(longest, _display_width(str(val)))
        ws.column_dimensions[letter].width = min(max(longest + padding, min_width), max_width)


def _has_data_rows(path: Path) -> bool:
    wb = load_workbook(path, read_only=True)
    try:
        return find_last_data_row(wb.active) > 1
    finally:
        wb.close()


def write_output(
    output_path: Path,
    new_rows: list[OutputRow],
    rules: Rules,
    template_headers: list[str],
    append: bool = True,
) -> tuple[int, list[str]]:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    existing_order_nos: set[str] = set()
    if append and output_path.exists():
        existing_order_nos = load_existing_order_nos(output_path)

    rows_to_write, skipped = filter_new_rows(new_rows, existing_order_nos)

    max_products = max((len(r.products) for r in rows_to_write), default=0)
    if append and output_path.exists():
        wb_old = load_workbook(output_path, read_only=True)
        old_headers = [
            wb_old.active.cell(1, c).value for c in range(1, wb_old.active.max_column + 1)
        ]
        wb_old.close()
        max_products = max(max_products, count_product_slots_in_headers(old_headers))

    headers = build_headers(max_products, template_headers, rules)

    wb = Workbook()
    ws = wb.active
    ws.title = rules.raw["global"]["sheet"]

    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)

    out_row_idx = 2

    if append and output_path.exists() and _has_data_rows(output_path):
        wb_old = load_workbook(output_path, read_only=True, data_only=True)
        ws_old = wb_old.active
        old_headers = [ws_old.cell(1, c).value for c in range(1, ws_old.max_column + 1)]
        for r in range(2, ws_old.max_row + 1):
            order_val = None
            for i, h in enumerate(old_headers):
                if h and str(h).startswith("用户订单号"):
                    order_val = ws_old.cell(r, i + 1).value
                    break
            if order_val is None or str(order_val).strip() == "":
                continue
            old_vals = {
                old_headers[i]: ws_old.cell(r, i + 1).value for i in range(len(old_headers))
            }
            new_vals = [old_vals.get(h, "") for h in headers]
            for c, v in enumerate(new_vals, 1):
                ws.cell(out_row_idx, c, v)
            out_row_idx += 1
        wb_old.close()

    for row in rows_to_write:
        vals = row_to_values(row, headers, rules)
        for c, v in enumerate(vals, 1):
            ws.cell(out_row_idx, c, v)
        out_row_idx += 1

    autofit_column_widths(ws)
    wb.save(output_path)
    return len(rows_to_write), skipped
