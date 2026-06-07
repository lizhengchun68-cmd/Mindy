from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from src.common.models import OrderRow

INPUT_COLUMNS = {
    "订单号": "order_no",
    "款式": "style",
    "数量": "quantity",
    "联系方式": "contact",
    "收件人名称": "recipient",
    "收件人国家": "country",
    "城市": "city",
    "州": "state",
    "邮编": "zip_code",
    "详细地址": "address",
    "备注": "remark",
}


def list_input_files(input_dir: Path) -> list[Path]:
    files = []
    for p in sorted(input_dir.glob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        files.append(p)
    return files


def read_orders(path: Path, sheet: str = "Sheet1") -> list[OrderRow]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{path} 缺少工作表 {sheet}")
    ws = wb[sheet]

    headers: dict[int, str] = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(1, c).value
        if val:
            headers[c] = str(val).strip()

    rows: list[OrderRow] = []
    for r in range(2, ws.max_row + 1):
        data: dict[str, object] = {}
        for c, name in headers.items():
            if name in INPUT_COLUMNS:
                data[INPUT_COLUMNS[name]] = ws.cell(r, c).value

        order_no = data.get("order_no")
        if order_no is None or str(order_no).strip() == "":
            continue

        rows.append(
            OrderRow(
                order_no=str(order_no).strip(),
                style=str(data.get("style") or "").strip(),
                quantity=data.get("quantity"),
                contact=_str_or_none(data.get("contact")),
                recipient=_str_or_none(data.get("recipient")),
                country=_str_or_none(data.get("country")),
                city=_str_or_none(data.get("city")),
                state=_str_or_none(data.get("state")),
                zip_code=_str_or_none(data.get("zip_code")),
                address=_str_or_none(data.get("address")),
                remark=_str_or_none(data.get("remark")),
                source_row=r,
            )
        )
    wb.close()
    return rows


def group_by_order(rows: list[OrderRow]) -> list[list[OrderRow]]:
    groups: dict[str, list[OrderRow]] = defaultdict(list)
    for row in rows:
        groups[row.order_no].append(row)
    result = []
    for order_no in sorted(groups.keys(), key=lambda k: min(r.source_row for r in groups[k])):
        result.append(sorted(groups[order_no], key=lambda r: r.source_row))
    return result


def _str_or_none(val: object) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None
